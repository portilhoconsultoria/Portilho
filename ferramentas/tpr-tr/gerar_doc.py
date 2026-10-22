"""Relatório TPR/TR -> PDF timbrado + Word, a partir do markdown de relatorio_tpr_tr.py.
Uso: python3 -I gerar_doc.py <relatorio.md> <saida_sem_extensao>"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from timbrado_engine import gerar_pdf
from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL, CINZA, TXT = "1C6086", "5F6F74", "1D2A2E"
DIAS = ["segunda-feira", "terça-feira", "quarta-feira", "quinta-feira", "sexta-feira", "sábado", "domingo"]


def preparar(md):
    # os símbolos de meta não existem na fonte do timbrado: viram texto
    md = md.replace(" ✅", " (na meta)").replace(" ❌", " (fora)")
    m = re.search(r"— (\d{4})-(\d{2})-(\d{2})$", md.splitlines()[0])
    if m:
        import datetime
        d = datetime.date(int(m[1]), int(m[2]), int(m[3]))
        linhas = md.splitlines()
        linhas[0] = f"# Tempo de resposta do comercial — {d:%d/%m/%Y}"
        linhas.insert(1, f"### Floricultura Verdes Plantas · {DIAS[d.weekday()]}, {d:%d/%m/%Y} · TPR e TR medidos mensagem a mensagem no DKW")
        md = "\n".join(linhas)
    return md


def sombrear(cel, cor):
    tcPr = cel._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), cor)
    tcPr.append(sh)


def run(p, texto, bold=False, italic=False, size=10.5, cor=TXT):
    partes = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto)
    for parte in partes:
        if not parte:
            continue
        b = bold or parte.startswith("**")
        t = parte.strip("*`") if parte.startswith(("**", "`")) else parte
        r = p.add_run(t); r.bold = b; r.italic = italic; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(cor)
        r.font.name = "Calibri"


def gerar_docx(md, caminho):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = Mm(30); sec.bottom_margin = Mm(20)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    # cabeçalho: faixa azul com a marca
    hp = sec.header.paragraphs[0]
    pPr = hp._p.get_or_add_pPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), AZUL); pPr.append(sh)
    r = hp.add_run("  PORTILHO "); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RGBColor(255, 255, 255)
    r = hp.add_run("Inteligência Comercial"); r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(255, 255, 255)
    fp = sec.footer.paragraphs[0]; run(fp, "Portilho Inteligência Comercial", size=8.5, cor=CINZA)
    linhas = md.splitlines(); i = 0
    while i < len(linhas):
        l = linhas[i].rstrip()
        if not l.strip():
            i += 1; continue
        if l.startswith("|") and i + 1 < len(linhas) and re.match(r"^\s*\|[\s:|-]+\|\s*$", linhas[i + 1]):
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                bloco.append(linhas[i]); i += 1
            cab = [c.strip() for c in bloco[0].strip().strip("|").split("|")]
            corpo = [[c.strip() for c in b.strip().strip("|").split("|")] for b in bloco[2:]]
            t = doc.add_table(rows=1 + len(corpo), cols=len(cab)); t.style = "Table Grid"
            for j, c in enumerate(cab):
                cel = t.rows[0].cells[j]; sombrear(cel, AZUL); cel.paragraphs[0].text = ""
                run(cel.paragraphs[0], c, bold=True, size=9, cor="FFFFFF")
            for k, linha in enumerate(corpo, 1):
                for j, c in enumerate(linha[:len(cab)]):
                    cel = t.rows[k].cells[j]
                    if k % 2 == 0: sombrear(cel, "F4F6F6")
                    run(cel.paragraphs[0], c, size=9)
            doc.add_paragraph()
            continue
        if l.startswith("# "):
            p = doc.add_paragraph(); run(p, l[2:].upper(), bold=True, size=14, cor=AZUL)
        elif l.startswith("### "):
            p = doc.add_paragraph(); run(p, l[4:], italic=True, size=10, cor=CINZA)
        elif l.startswith("## "):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); run(p, l[3:], bold=True, size=11.5, cor=AZUL)
        elif l.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); run(p, l[2:])
        else:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; run(p, l)
        i += 1
    doc.save(caminho)
    return caminho


if __name__ == "__main__":
    md = preparar(open(sys.argv[1], encoding="utf-8").read())
    base = sys.argv[2]
    print(gerar_pdf(base + ".pdf", md, marca="portilho"))
    print(gerar_docx(md, base + ".docx"))
